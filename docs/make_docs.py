"""
경진대회 제출용 Word 문서 4종 생성 스크립트
실행: python docs/make_docs.py
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path

DOCS_DIR = Path(__file__).resolve().parent

# ── 공통 헬퍼 ────────────────────────────────────────────────────────────────
RED   = RGBColor(0xD6, 0x27, 0x28)
GRAY  = RGBColor(0x55, 0x55, 0x55)
BLACK = RGBColor(0x11, 0x11, 0x11)

def _set_font(run, size=11, bold=False, color=None, name="맑은 고딕"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    r = run._r
    rPr = r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:eastAsia'), name)
    rPr.insert(0, rFonts)

def heading1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after  = Pt(6)
    run = p.add_run(text)
    _set_font(run, size=16, bold=True, color=RED)
    return p

def heading2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    _set_font(run, size=13, bold=True, color=BLACK)
    return p

def heading3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text)
    _set_font(run, size=11, bold=True, color=BLACK)
    return p

def body(doc, text, indent=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.left_indent = Cm(0.7)
    run = p.add_run(text)
    _set_font(run, size=10.5, color=GRAY)
    return p

def bold_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    _set_font(run, size=10.5, bold=True)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.7)
    run = p.add_run(text)
    _set_font(run, size=10.5, color=GRAY)
    return p

def code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Cm(0.7)
    p.paragraph_format.right_indent = Cm(0.7)
    p.paragraph_format.space_after  = Pt(6)
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    # 배경색 (연회색)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F5F5F5')
    pPr.append(shd)
    return p

def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 헤더 행
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in hdr_cells[i].paragraphs[0].runs:
            _set_font(run, size=10, bold=True)
        # 헤더 배경색
        tc = hdr_cells[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'),   'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'),  'D62728')
        tcPr.append(shd)
        for run in hdr_cells[i].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # 데이터 행
    for r_i, row_data in enumerate(rows):
        cells = table.rows[r_i+1].cells
        for c_i, val in enumerate(row_data):
            cells[c_i].text = str(val)
            cells[c_i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in cells[c_i].paragraphs[0].runs:
                _set_font(run, size=10)
        # 짝수 행 배경
        if r_i % 2 == 1:
            for cell in cells:
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'),   'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'),  'FFF5F5')
                tcPr.append(shd)

    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    return table

def title_page(doc, title, subtitle="2026년 제4회 전북 청년 AI·빅데이터 경진대회"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(60)
    p.paragraph_format.space_after  = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    _set_font(run, size=22, bold=True, color=RED)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(4)
    run2 = p2.add_run(subtitle)
    _set_font(run2, size=11, color=GRAY)

    doc.add_paragraph()
    doc.add_page_break()

def set_margins(doc, top=2.5, bottom=2.5, left=3.0, right=3.0):
    section = doc.sections[0]
    section.top_margin    = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin   = Cm(left)
    section.right_margin  = Cm(right)


# ════════════════════════════════════════════════════════════════════════════
# ① 최종 결과보고서
# ════════════════════════════════════════════════════════════════════════════
def make_report():
    doc = Document()
    set_margins(doc)
    title_page(doc, "최종 결과보고서", "전북 전기화재 예방 우선지원 모델\n2026년 제4회 전북 청년 AI·빅데이터 경진대회")

    heading1(doc, "1. 프로젝트 개요")
    heading2(doc, "1.1 분석 목적")
    body(doc, "전북특별자치도 14개 시군의 전기화재 데이터를 분석하여, 한정된 예방 자원(인력·장비)을 어느 지역에 먼저 투입해야 하는지 데이터 기반으로 결정하는 우선지원 모델을 구축한다.")

    heading2(doc, "1.2 분석 배경")
    bullet(doc, "전북 인구 10만명당 전기화재 건수는 전국 상위권 수준")
    bullet(doc, "전기안전 정기점검 부적합률은 전국 평균 수준임에도 화재율이 높음")
    bullet(doc, "기존 점검 체계가 고령 1인가구 단독주택에 충분히 닿지 못하고 있다는 구조적 문제 존재")

    heading2(doc, "1.3 분석 범위")
    add_table(doc,
        ["항목", "내용"],
        [["대상 지역", "전북특별자치도 14개 시군"],
         ["분석 기간", "2021~2025년 (5개년)"],
         ["분석 도구", "Python 3.13, Streamlit, scikit-learn, statsmodels, Anthropic Claude API"]],
        col_widths=[4, 12]
    )

    doc.add_paragraph()
    heading1(doc, "2. 데이터 출처")
    add_table(doc,
        ["데이터", "출처", "기간"],
        [["전기화재 시군별 발생 현황", "한국전기안전공사(KESCO)", "2021~2025"],
         ["시도별 화재 발생 현황", "한국전기안전공사(KESCO)", "2022~2024"],
         ["전기화재 발화원인별 통계", "한국전기안전공사(KESCO)", "2022~2025"],
         ["정기점검 행정구역별 실시현황", "한국전기안전공사(KESCO)", "2016~2025"],
         ["1인가구 현황 (연령별)", "통계청 KOSIS", "2020~2025"],
         ["장래인구추계", "통계청", "2052년 중위 시나리오"],
         ["노후주택 현황", "국토교통부", "2023"],
         ["비상소화장치 현황", "전북특별자치도 소방본부", "2024"],
         ["전국 소방서 좌표", "소방청", "2024"]],
        col_widths=[6, 6, 4]
    )

    doc.add_paragraph()
    heading1(doc, "3. 분석 방법론")
    heading2(doc, "3.1 종합위험점수 산출")
    body(doc, "화재율·재산피해액·고령1인가구밀도 3개 지표의 순위를 합산하여 시군별 종합위험점수를 산출하고 우선순위를 부여했다.")
    code_block(doc, "종합위험점수 = rank(천호당화재율) + rank(평균재산피해액) + rank(고령1인가구밀도)")

    heading2(doc, "3.2 상관관계 분석 — 4단계 가설 검증")
    add_table(doc,
        ["가설", "변수", "상관계수(r)", "결과"],
        [["가설 1", "노후주택비율 전체", "0.152", "기각 (아파트 포함 희석)"],
         ["가설 2", "노후단독주택비율", "0.37", "부분 지지"],
         ["가설 3", "고령1인가구밀도", "0.801", "강한 상관"],
         ["가설 4 ✅", "고령단독주택거주율", "0.854", "최강 설명력 채택"]],
        col_widths=[2.5, 5, 3, 5.5]
    )

    heading2(doc, "3.3 회귀분석 (OLS)")
    body(doc, "최종 모형: 천호당 화재율 ~ 고령 1인가구 단독주택 거주율")
    add_table(doc,
        ["통계량", "값"],
        [["상관계수 (r)", "0.854"],
         ["결정계수 (R²)", "0.730"],
         ["p-value", "0.0001"],
         ["표본", "14개 시군"]],
        col_widths=[6, 10]
    )
    body(doc, "고령 1인가구 단독주택 거주율이 1%p 높아질 때마다 천호당 전기화재 건수가 통계적으로 유의하게 증가한다.", indent=True)
    body(doc, "다중공선성 검증: 고령1인가구밀도 + 노후주택비율 동시 투입 시, 노후주택비율은 p=0.09로 유의하지 않음 → 허위상관 확인.", indent=True)

    heading2(doc, "3.4 K-means 클러스터링 (비지도학습)")
    bullet(doc, "특성: 고령단독주택거주율, 천호당화재율, 고령1인가구밀도 (3개)")
    bullet(doc, "클러스터 수: k=3")
    bullet(doc, "실루엣 점수: 0.45 (0.4 이상 = 유의미한 분리)")
    bullet(doc, "레이블 매핑: 천호당화재율 평균 내림차순으로 고위험·중위험·관리 등급 고정")

    heading2(doc, "3.5 생성형 AI 브리핑 (Anthropic Claude claude-sonnet-4-6)")
    body(doc, "시군별 데이터(ML 등급, 자원배분 갭, 화재율, 고령 지표)를 포함한 구조화 프롬프트를 사용하여 진단 요약·우선 개입 3가지·근거를 포함한 브리핑을 스트리밍으로 생성한다.")

    doc.add_paragraph()
    heading1(doc, "4. 핵심 발견")
    heading2(doc, "4.1 기존 점검 체계의 사각지대")
    body(doc, "전북 전기안전 정기점검 부적합률은 전국 평균 수준이지만 화재율은 전국 상위권. 기존 점검(주로 상업·사업용 설비 대상)이 고령 1인가구 단독주택에 닿지 않고 있음을 시사한다.")

    heading2(doc, "4.2 단독주택 거주 구조가 핵심 위험 요인")
    body(doc, "아파트는 관리사무소가 배선을 점검하지만 단독주택은 거주자 본인이 관리해야 한다. 고령 1인가구가 단독주택에 집중적으로 거주하면서 배선 관리가 방치되는 것이 화재로 이어지는 핵심 메커니즘이다.")

    heading2(doc, "4.3 위험과 자원의 불일치")
    body(doc, "K-means 고위험 등급이면서 비상소화장치 밀도가 전북 중앙값 미만인 시군 존재 → 위험도와 인프라 배치가 불일치한다.")

    heading2(doc, "4.4 우선순위 상위 시군")
    add_table(doc,
        ["순위", "시군", "천호당화재율", "ML 등급"],
        [["1", "진안군", "1.152건", "고위험"],
         ["1", "남원시", "0.899건", "고위험"],
         ["3", "임실군", "1.155건", "고위험"],
         ["4", "김제시", "0.800건", "중위험"]],
        col_widths=[2, 4, 5, 5]
    )

    doc.add_paragraph()
    heading1(doc, "5. 예방 전략")
    heading2(doc, "POLICY 01 — 찾아가는 전기안전 점검 우선 배정")
    body(doc, "우선순위 상위 시군의 고령 1인가구 단독주택을 대상으로 한국전기안전공사 무료 방문 점검을 우선 배정. 경로당·복지관 연계 신청 접수 체계 구축.")

    heading2(doc, "POLICY 02 — 비상소화장치 갭 지역 우선 보강")
    body(doc, "ML 고위험 + 장치밀도 부족 지역에 비상소화장치 추가 배치. 소방관서가 1~2개인 진안·임실·무주·순창은 대응 시간 보완책 병행 필요.")

    heading2(doc, "POLICY 03 — 노후 배선 교체비 지원 확대")
    body(doc, "전기화재 주원인(절연열화·트래킹)은 노후 배선 교체로 직접 예방 가능. 저소득 고령 1인가구 대상 배선·누전차단기 교체비 지원사업을 우선순위 상위 시군에 집중 배정.")

    doc.add_paragraph()
    heading1(doc, "6. 분석의 한계")
    bullet(doc, "표본이 14개 시군으로 소수 → 신뢰구간이 넓고, 읍면동 단위 세분화 데이터로 재검증 필요")
    bullet(doc, "고령단독주택거주율과 고령1인가구밀도 간 다중공선성(r≈0.85) 존재")
    bullet(doc, "상관관계·회귀분석이며 직접적인 인과관계 증명은 아님")
    bullet(doc, "K-means 클러스터 경계는 패턴 참고용 (n=14 소표본)")
    bullet(doc, "정책 효과는 사업 시행 후 데이터로 재보정 필요")

    doc.save(DOCS_DIR / "01_최종결과보고서.docx")
    print("✓ 01_최종결과보고서.docx")


# ════════════════════════════════════════════════════════════════════════════
# ② 데이터셋 설명서
# ════════════════════════════════════════════════════════════════════════════
def make_dataset():
    doc = Document()
    set_margins(doc)
    title_page(doc, "데이터셋 설명서", "전북 전기화재 예방 우선지원 모델")

    heading1(doc, "1. 데이터 파일 목록")
    add_table(doc,
        ["파일명", "출처", "기간", "인코딩"],
        [["fire_raw.csv", "한국전기안전공사", "2021~2025", "UTF-8"],
         ["summary.csv", "자체 산출", "2021~2025 평균", "UTF-8"],
         ["전기화재_지역별.csv", "한국전기안전공사", "2021~2025", "CP949"],
         ["fire_cause_by_year.csv", "한국전기안전공사", "2022~2025", "CP949"],
         ["sido_fire.csv", "한국전기안전공사", "2022~2024", "CP949"],
         ["정기점검_행정구역별.csv", "한국전기안전공사", "2016~2025", "CP949"],
         ["노후주택현황.csv", "국토교통부", "2023", "CP949"],
         ["비상소화장치현황.csv", "전북도 소방본부", "2024", "CP949"],
         ["전국소방서_좌표.csv", "소방청", "2024", "CP949"],
         ["kosis_elderly_single_household.csv", "통계청 KOSIS", "2020~2025", "CP949"],
         ["population_projection_2052.csv", "통계청", "2052 중위", "CP949"]],
        col_widths=[6, 5, 3.5, 2.5]
    )

    doc.add_paragraph()
    heading1(doc, "2. 파일별 상세 설명")

    heading2(doc, "2.1 fire_raw.csv")
    body(doc, "전북 14개 시군의 연도별 전기화재 건수 및 재산피해 (분석 핵심 원본).")
    add_table(doc,
        ["컬럼", "설명"],
        [["시군", "전북 14개 시군명"],
         ["연도", "2021~2025"],
         ["화재건수", "해당 연도 전기화재 발생 건수"],
         ["재산피해_천원", "해당 연도 재산피해액 (단위: 천원)"]],
        col_widths=[5, 11]
    )
    body(doc, "※ 전처리: 2021년 전주시는 완산구·덕진구로 분리 기재 → 합산하여 전주시로 통합.", indent=True)

    heading2(doc, "2.2 summary.csv")
    body(doc, "14개 시군의 5개년(2021~2025) 평균값 기반 종합 지표. 분석의 핵심 데이터프레임.")
    add_table(doc,
        ["컬럼", "설명", "단위"],
        [["시군", "전북 14개 시군", "-"],
         ["천호당화재율", "평균화재건수 / 총주택수 × 1000", "건/천호"],
         ["평균피해액_천원", "5개년 평균 재산피해액", "천원"],
         ["노후주택비율", "30년 이상 노후주택 비율", "%"],
         ["고령1인가구밀도", "65세+ 1인가구 수 / 총주택수 × 1000", "명/천호"],
         ["고령단독주택거주율", "고령 1인가구 중 단독주택 거주 비율", "%"],
         ["종합위험점수", "rank 3개 합산 (낮을수록 시급)", "점"],
         ["우선순위", "종합위험점수 순위", "1~14"]],
        col_widths=[5, 8, 3]
    )

    heading2(doc, "2.3 fire_cause_by_year.csv")
    body(doc, "전국 전기화재 발화원인별 연도 통계. 회귀분석 배경 자료.")
    add_table(doc,
        ["컬럼", "설명"],
        [["연도", "2022~2025"],
         ["발화 원인", "절연열화에 의한 단락, 트래킹 등"],
         ["화재 건수", "건"],
         ["점유율", "전체 전기화재 대비 비율 (%)"]],
        col_widths=[5, 11]
    )

    heading2(doc, "2.4 정기점검_행정구역별.csv")
    body(doc, "한국전기안전공사 정기점검 실시 현황. 역설 분석(부적합률 평균 vs 높은 화재율)에 사용.")
    add_table(doc,
        ["컬럼", "설명"],
        [["행정구역별", "시도명"],
         ["연도", "2016~2025 (2019년 기준변경 제외)"],
         ["부적합률", "정기점검 중 부적합 판정 비율 (%)"]],
        col_widths=[5, 11]
    )
    body(doc, "※ 전처리: 2019년은 전국 기준 변경으로 이상치 발생 → 분석에서 제외.", indent=True)

    heading2(doc, "2.5 비상소화장치현황.csv")
    body(doc, "전북특별자치도 비상소화장치 설치 위치 및 좌표. 지도 시각화 및 갭 분석에 사용.")
    body(doc, "※ 전처리: 설치위치에서 정규식(전북특별자치도 \\S+시|\\S+군)으로 시군명 추출.", indent=True)

    heading2(doc, "2.6 kosis_elderly_single_household.csv")
    body(doc, "통계청 KOSIS 연령별·행정구역별 1인가구 수. 고령화 추이 분석에 사용.")
    body(doc, "※ 전처리: 연도 컬럼 쉼표 제거 후 pd.to_numeric() 변환 필수 (문자열 → 숫자).", indent=True)

    doc.add_paragraph()
    heading1(doc, "3. 데이터 전처리 요약")
    add_table(doc,
        ["이슈", "해결 방법"],
        [["2021년 전주시 완산구·덕진구 분리 기재", "두 구 합산 후 전주시로 통합"],
         ["2019년 정기점검 기준 변경으로 이상치", "해당 연도 분석 제외"],
         ["kosis 파일 숫자에 쉼표 포함 (문자열)", "str.replace(',', '') → pd.to_numeric 변환"],
         ["비상소화장치 시군 정보 없음", "설치위치 주소에서 정규식으로 추출"]],
        col_widths=[8, 8]
    )

    doc.save(DOCS_DIR / "02_데이터셋설명서.docx")
    print("✓ 02_데이터셋설명서.docx")


# ════════════════════════════════════════════════════════════════════════════
# ③ 코드·분석파일 설명서
# ════════════════════════════════════════════════════════════════════════════
def make_code():
    doc = Document()
    set_margins(doc)
    title_page(doc, "코드·분석파일 설명서", "전북 전기화재 예방 우선지원 모델")

    heading1(doc, "1. 프로젝트 디렉토리 구조")
    code_block(doc,
        "전북_전기화재_대시보드/\n"
        "├── app.py                  # 메인 Streamlit 애플리케이션\n"
        "├── requirements.txt        # Python 패키지 의존성\n"
        "├── .env                    # API 키 설정 (git 제외)\n"
        "├── .env.example            # API 키 설정 예시\n"
        "├── data/                   # 분석 데이터 (CSV)\n"
        "└── docs/                   # 제출 문서"
    )

    heading1(doc, "2. app.py 구조")
    heading2(doc, "2.1 섹션별 구성")
    add_table(doc,
        ["구역", "라인(약)", "내용"],
        [["임포트 및 초기화", "1~17", "패키지 임포트, dotenv 로드"],
         ["페이지 설정", "20", "set_page_config (🚒 아이콘, wide 레이아웃)"],
         ["데이터 로드", "28~120", "CSV 로드, 상관계수 계산, 노후단독주택비율 산출"],
         ["K-means 클러스터링", "122~143", "비지도학습 위험 등급 분류"],
         ["Folium 지도 함수", "146~244", "render_folium_map()"],
         ["AI 브리핑 함수", "247~337", "call_llm(), template_briefing()"],
         ["사이드바", "340~365", "앵커 네비게이션"],
         ["CSS 주입", "368~398", "커스텀 스타일"],
         ["상단 KPI 카드", "405~417", "4개 핵심 지표"],
         ["섹션 1 — 전국 비교", "418~800", "시도별 화재율, 추이, 역설, 고령화"],
         ["섹션 2 — 연도별 추이", "801~812", "시군별 화재 건수 추이"],
         ["섹션 3 — 위험지수 순위", "813~834", "막대 차트 + 테이블"],
         ["섹션 4 — 상관관계", "835~908", "4가설 검증 산점도"],
         ["섹션 5 — 회귀분석", "909~1050", "3단계 OLS 검증"],
         ["섹션 6 — ML 위험 분류", "1051~1101", "K-means 버블 차트"],
         ["섹션 7 — 지도", "1102~1163", "Folium 지도 + 갭 산점도"],
         ["섹션 8 — 시뮬레이터", "1164~1261", "자원배분 시뮬레이션 + 갭 분석"],
         ["섹션 9 — AI 브리핑", "1262~1324", "시군 선택 + 스트리밍 브리핑"],
         ["섹션 10 — 예방 전략", "1325~끝", "핵심 발견 + 정책 제언 카드"]],
        col_widths=[5, 2.5, 8.5]
    )

    doc.add_paragraph()
    heading1(doc, "3. 주요 함수 설명")
    heading2(doc, "3.1 데이터 로드")
    add_table(doc,
        ["함수", "반환 타입", "설명"],
        [["load_data()", "DataFrame × 2", "fire_raw.csv, summary.csv 로드 (캐시)"],
         ["load_fire_devices()", "DataFrame", "비상소화장치 현황 + 시군명 추출"],
         ["load_fire_stations()", "DataFrame", "전국 소방서 좌표 → 전북 필터링"],
         ["load_cause_data()", "DataFrame", "전기화재 발화원인 연도 통계"],
         ["load_old_housing()", "DataFrame", "노후주택 현황 (전주시 합산)"]],
        col_widths=[5, 3.5, 7.5]
    )

    heading2(doc, "3.2 분석 함수")
    add_table(doc,
        ["함수", "반환 타입", "설명"],
        [["fit_regression(df)", "OLSResults × 3", "3단계 OLS 회귀 적합 (cache_resource)"],
         ["run_kmeans(df)", "(DataFrame, float)", "K-means 분류 + 실루엣 점수 반환"]],
        col_widths=[5, 3.5, 7.5]
    )

    heading2(doc, "3.3 시각화 함수")
    add_table(doc,
        ["함수", "반환 타입", "설명"],
        [["render_folium_map(...)", "folium.Map", "전북 지도 (위험 원, 소화장치, 소방서)"]],
        col_widths=[5, 3.5, 7.5]
    )

    heading2(doc, "3.4 AI 브리핑 함수")
    add_table(doc,
        ["함수", "반환 타입", "설명"],
        [["call_llm(region, row, api_key)", "Generator", "Claude API 스트리밍 브리핑"],
         ["template_briefing(region, row)", "str", "API 키 미설정 시 규칙 기반 대체"]],
        col_widths=[5, 3.5, 7.5]
    )

    doc.add_paragraph()
    heading1(doc, "4. 내러티브 흐름")
    add_table(doc,
        ["순서", "섹션", "역할"],
        [["①", "전국 비교", "문제 제기: 전북이 전국 상위권"],
         ["②", "연도별 추이", "최근 5년 트렌드"],
         ["③", "위험지수 순위", "전북 내 어디가 위험한가"],
         ["④", "상관관계", "왜 위험한가 — 가설 검증"],
         ["⑤", "회귀분석", "통계적 증명"],
         ["⑥", "ML 위험 분류", "구조적 그룹 분류"],
         ["⑦", "지도", "종합 시각화"],
         ["⑧", "시뮬레이터", "자원을 어디에 배분할 것인가"],
         ["⑨", "AI 브리핑", "시군별 맞춤 AI 전략"],
         ["⑩", "예방 전략", "결론 및 정책 제언"]],
        col_widths=[1.5, 4, 10.5]
    )

    doc.add_paragraph()
    heading1(doc, "5. 핵심 변수 목록")
    add_table(doc,
        ["변수", "타입", "설명"],
        [["fire_df", "DataFrame", "연도별 시군 화재 건수"],
         ["summary_df", "DataFrame", "14개 시군 종합 지표 (cluster_label 포함)"],
         ["device_df", "DataFrame", "비상소화장치 현황"],
         ["station_df", "DataFrame", "전북 소방서 좌표"],
         ["r_single", "float", "고령단독주택거주율 상관계수 (동적 계산)"],
         ["sil_score", "float", "K-means 실루엣 점수"],
         ["reg_new", "OLSResults", "최종 회귀모형 (고령단독주택거주율)"],
         ["gap_df", "DataFrame", "비상소화장치 갭 분석 결과 (섹션 8에서 정의)"],
         ["jeonbuk_rank", "int", "전북 전국 화재율 순위 (섹션 1에서 정의)"],
         ["EFFECTIVE_ANTHROPIC_KEY", "str", "환경변수에서 로드한 API 키"]],
        col_widths=[5, 3, 8]
    )

    doc.save(DOCS_DIR / "03_코드분석파일설명서.docx")
    print("✓ 03_코드분석파일설명서.docx")


# ════════════════════════════════════════════════════════════════════════════
# ④ 실행환경 설명서
# ════════════════════════════════════════════════════════════════════════════
def make_env():
    doc = Document()
    set_margins(doc)
    title_page(doc, "실행환경 설명서", "전북 전기화재 예방 우선지원 모델")

    heading1(doc, "1. 시스템 요구사항")
    add_table(doc,
        ["항목", "요구사항"],
        [["OS", "macOS 12+, Windows 10+, Ubuntu 20.04+"],
         ["Python", "3.10 이상"],
         ["메모리", "4GB RAM 이상"],
         ["인터넷 연결", "AI 브리핑 기능 사용 시 필요"]],
        col_widths=[5, 11]
    )

    doc.add_paragraph()
    heading1(doc, "2. Python 패키지 목록")
    add_table(doc,
        ["패키지", "버전", "역할"],
        [["streamlit", "≥1.30", "웹 대시보드 프레임워크"],
         ["pandas", "≥2.0", "데이터 로드 및 처리"],
         ["plotly", "≥5.18", "인터랙티브 차트"],
         ["statsmodels", "≥0.14", "OLS 회귀분석"],
         ["scikit-learn", "≥1.4", "K-means 클러스터링, StandardScaler"],
         ["folium", "≥0.15", "지도 시각화 (OpenStreetMap)"],
         ["streamlit-folium", "≥0.20", "Streamlit에 Folium 지도 삽입"],
         ["anthropic", "≥0.25", "Claude API 연동 (AI 브리핑)"],
         ["python-dotenv", "≥1.0", ".env 파일에서 환경변수 로드"]],
        col_widths=[4, 3, 9]
    )

    doc.add_paragraph()
    heading1(doc, "3. 설치 방법")
    heading2(doc, "3.1 가상환경 생성 (권장)")
    code_block(doc,
        "# macOS / Linux\n"
        "python3 -m venv venv\n"
        "source venv/bin/activate\n\n"
        "# Windows\n"
        "python -m venv venv\n"
        "venv\\Scripts\\activate"
    )
    heading2(doc, "3.2 패키지 설치")
    code_block(doc, "pip install -r requirements.txt")

    doc.add_paragraph()
    heading1(doc, "4. 환경변수 설정 (API 키)")
    body(doc, "AI 브리핑 기능을 사용하려면 Anthropic API 키가 필요합니다.")
    heading2(doc, "4.1 .env 파일 생성")
    code_block(doc, "cp .env.example .env")
    heading2(doc, "4.2 API 키 입력")
    body(doc, ".env 파일을 텍스트 편집기로 열고 API 키를 입력합니다.")
    code_block(doc, "ANTHROPIC_API_KEY=sk-ant-여기에_실제_키를_입력하세요")
    body(doc, "※ .env 파일은 .gitignore에 포함되어 git에 업로드되지 않습니다.", indent=True)
    body(doc, "※ API 키가 없어도 앱은 정상 실행됩니다. AI 브리핑 클릭 시 템플릿 브리핑으로 대체됩니다.", indent=True)

    doc.add_paragraph()
    heading1(doc, "5. 앱 실행")
    code_block(doc, "streamlit run app.py")
    body(doc, "실행 후 브라우저에서 자동으로 http://localhost:8501 이 열립니다.")
    body(doc, "※ app.py가 있는 프로젝트 루트 디렉토리에서 실행해야 합니다.", indent=True)
    heading2(doc, "포트 변경이 필요한 경우")
    code_block(doc, "streamlit run app.py --server.port 8502")

    doc.add_paragraph()
    heading1(doc, "6. 자주 발생하는 오류 및 해결")
    add_table(doc,
        ["오류", "원인", "해결 방법"],
        [["ModuleNotFoundError", "패키지 미설치", "pip install -r requirements.txt"],
         ["FileNotFoundError: data/...", "실행 위치 오류", "app.py 디렉토리에서 실행"],
         ["AI 브리핑 오류", "API 키 미설정 또는 만료", ".env에서 API 키 확인"],
         ["streamlit_folium 없음", "패키지 미설치", "pip install streamlit-folium"]],
        col_widths=[4.5, 4.5, 7]
    )

    doc.add_paragraph()
    heading1(doc, "7. 실행 확인 체크리스트")
    bullet(doc, "브라우저에서 http://localhost:8501 접속")
    bullet(doc, "페이지 타이틀 \"전북 전기화재 예방 우선지원 모델\" 확인")
    bullet(doc, "브라우저 탭 아이콘 🚒 확인")
    bullet(doc, "사이드바에 \"전북 화재 예방\" 표시 확인")
    bullet(doc, "전국 비교 섹션 차트 렌더링 확인")
    bullet(doc, "지도 섹션에서 전북 14개 시군 원 표시 확인")
    bullet(doc, "AI 브리핑 — \"AI 브리핑 생성\" 버튼 클릭 시 스트리밍 출력 확인")

    doc.save(DOCS_DIR / "04_실행환경설명서.docx")
    print("✓ 04_실행환경설명서.docx")


if __name__ == "__main__":
    make_report()
    make_dataset()
    make_code()
    make_env()
    print("\n완료: docs/ 폴더에 4개 Word 파일 생성됨")
